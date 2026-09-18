"""
utils/document_generator.py
============================
Generates supplementary DOCX documents — cover letter and ATS analysis report.

NOTE: The resume itself is no longer rebuilt from scratch.
      Resume output is handled by utils/inplace_editor.py, which only
      modifies the Technical Skills section of the original uploaded file.
      All resume-rebuild code (section rewriting, page-budget trimming,
      bullet ranking, experience-duration calculation, external subprocess
      calls) has been removed from this module.
"""

from docx import Document
from datetime import datetime
import io
import re


# ---------------------------------------------------------------------------
# Text sanitisation
# ---------------------------------------------------------------------------

def sanitize_text(text) -> str:
    """Remove invalid XML / control characters for DOCX compatibility."""
    if not text or not isinstance(text, str):
        return ""
    text = re.sub(r'[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F-\x9F]', '', text)
    text = text.replace('\x00', '').replace('\ufeff', '')
    return str(text).strip()


# ---------------------------------------------------------------------------
# Public — create_cover_letter_docx
# ---------------------------------------------------------------------------

def create_cover_letter_docx(
    resume_text: str,
    job_keywords: dict,
    score_data: dict,
    archetype_data: dict = None,
) -> io.BytesIO:
    """Generate a professional cover letter DOCX using only matched skills."""
    try:
        resume_text = sanitize_text(resume_text)

        # Extract name / contact from resume text
        lines = [l.strip() for l in resume_text.split('\n') if l.strip()]
        name = "Your Name"
        for line in lines[:10]:
            words = line.split()
            if 1 < len(words) <= 4 and '@' not in line and all(
                w[0].isupper() for w in words if w and w[0].isalpha()
            ):
                name = line
                break

        email_m = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', resume_text)
        email = email_m.group(0) if email_m else "your.email@example.com"

        phone_m = re.search(
            r'(\+91[-\s]?|0)?[6-9]\d{9}|(\+\d{1,3}[-\s]?)?\(?\d{3}\)?[-\s]?\d{3}[-\s]?\d{4}',
            resume_text
        )
        phone = phone_m.group(0) if phone_m else "+XX-XXXXXXXXXX"

        doc = Document()

        doc.add_paragraph(sanitize_text(name))
        doc.add_paragraph(f"{email} | {phone}")
        doc.add_paragraph(datetime.now().strftime("%B %d, %Y"))
        doc.add_paragraph()

        company  = sanitize_text(job_keywords.get("company", "Hiring Manager"))
        job_title = sanitize_text(job_keywords.get("job_title", "the position"))

        doc.add_paragraph(company)
        doc.add_paragraph(f"Re: Application for {job_title}")
        doc.add_paragraph()
        doc.add_paragraph("Dear Hiring Manager,")
        doc.add_paragraph()

        matched_skills = [sanitize_text(s) for s in score_data.get("matched_keywords", [])[:3] if s]
        skills_str = ", ".join(matched_skills) if matched_skills else "relevant technical skills"

        doc.add_paragraph(sanitize_text(
            f"I am writing to express my strong interest in the {job_title} position at {company}. "
            f"With proven expertise in {skills_str}, I am confident in my ability to make meaningful "
            f"contributions to your team from day one."
        ))
        doc.add_paragraph()

        doc.add_paragraph(sanitize_text(
            f"My background demonstrates a {score_data.get('ats_score', 0)}% keyword alignment with "
            f"your stated requirements, reflecting a strong match to the technical demands of this role."
        ))
        doc.add_paragraph()

        doc.add_paragraph(
            "I would welcome the opportunity to discuss how my experience aligns with your team's needs. "
            "Thank you for considering my application."
        )
        doc.add_paragraph()
        doc.add_paragraph("Sincerely,")
        doc.add_paragraph(sanitize_text(name))

        doc_bytes = io.BytesIO()
        doc.save(doc_bytes)
        doc_bytes.seek(0)
        return doc_bytes

    except Exception:
        doc = Document()
        doc.add_heading("Cover Letter", level=1)
        doc.add_paragraph("Dear Hiring Manager,")
        doc.add_paragraph("I am interested in this position.")
        doc.add_paragraph("Sincerely,")
        doc.add_paragraph("Your Name")
        doc_bytes = io.BytesIO()
        doc.save(doc_bytes)
        doc_bytes.seek(0)
        return doc_bytes


# ---------------------------------------------------------------------------
# Public — create_ats_report_docx
# ---------------------------------------------------------------------------

def create_ats_report_docx(
    score_data: dict,
    job_keywords: dict,
    archetype_data: dict = None,
    skill_buckets: dict = None,
) -> io.BytesIO:
    """
    Generate an ATS analysis report DOCX.

    Parameters
    ----------
    score_data    : output of ai_matcher.calculate_enhanced_ats_score()
    job_keywords  : output of parse_job_description()
    archetype_data: role archetype (optional)
    skill_buckets : three-bucket dict from classify_skills() (optional)
                    Keys: already_listed, optimizable, genuinely_missing
    """
    try:
        doc = Document()
        doc.add_heading("ATS ANALYSIS REPORT", level=1)
        doc.add_paragraph()

        if archetype_data:
            doc.add_heading("Detected Role", level=2)
            doc.add_paragraph(archetype_data.get("display_name", "General Technical Role"))
            doc.add_paragraph()

        doc.add_heading("ATS Score Breakdown", level=2)
        doc.add_paragraph(f"Overall ATS Score:    {score_data.get('ats_score', 0)}%")
        if "weighted_skill_score" in score_data:
            doc.add_paragraph(f"  Skill Match Score:  {score_data.get('weighted_skill_score', 0)}%   (weight: 60%)")
            doc.add_paragraph(f"  TF-IDF Score:       {score_data.get('tfidf_score', 0)}%   (weight: 40%)")
        doc.add_paragraph(f"Keywords Analysed:    {score_data.get('total_keywords', 0)}")
        doc.add_paragraph(f"Matched:              {score_data.get('matched_count', 0)}")
        doc.add_paragraph(f"Missing:              {score_data.get('missing_count', 0)}")
        doc.add_paragraph()

        if skill_buckets:
            # Already listed
            already = [sanitize_text(s) for s in skill_buckets.get("already_listed", []) if s]
            if already:
                doc.add_heading("✅ Already Listed in Skills Section", level=2)
                doc.add_paragraph(", ".join(already))
                doc.add_paragraph()

            # Optimizable (added to the output resume)
            optimizable = [sanitize_text(s) for s in skill_buckets.get("optimizable", []) if s]
            if optimizable:
                doc.add_heading("🔄 Skills Added to Resume (Evidenced Elsewhere)", level=2)
                doc.add_paragraph(
                    "These skills were found in your Experience/Projects sections and have been "
                    "appended to your Technical Skills section in the optimised resume:"
                )
                doc.add_paragraph(", ".join(optimizable))
                doc.add_paragraph()

            # Genuinely missing (advisory only — never added to resume)
            missing = [sanitize_text(s) for s in skill_buckets.get("genuinely_missing", []) if s]
            if missing:
                doc.add_heading("❌ Skills to Consider Adding (Advisory Only)", level=2)
                doc.add_paragraph(
                    "These skills appear in the job description but were NOT found anywhere in your "
                    "resume. They have NOT been added to any output document. Consider adding them "
                    "only if you genuinely possess them:"
                )
                doc.add_paragraph(", ".join(missing))
                doc.add_paragraph()
            elif not optimizable and not already:
                doc.add_heading("All Key Skills Matched", level=2)
                doc.add_paragraph("Your resume contains all required skills from the job description.")
                doc.add_paragraph()

        else:
            # Fallback: use raw matched/missing from score_data
            matched = [sanitize_text(s) for s in score_data.get("matched_keywords", []) if s]
            if matched:
                doc.add_heading("Matched Skills", level=2)
                doc.add_paragraph(", ".join(matched[:25]))
                doc.add_paragraph()

            missing = [sanitize_text(s) for s in score_data.get("missing_keywords", []) if s]
            if missing:
                doc.add_heading("Skills to Consider Adding (Advisory Only)", level=2)
                doc.add_paragraph(", ".join(missing[:25]))
                doc.add_paragraph()
            else:
                doc.add_heading("All Key Skills Matched", level=2)
                doc.add_paragraph("Your resume contains all required skills from the job description.")
                doc.add_paragraph()

        doc.add_heading("Job Details", level=2)
        doc.add_paragraph(f"Position: {sanitize_text(job_keywords.get('job_title', 'N/A'))}")
        doc.add_paragraph(f"Company:  {sanitize_text(job_keywords.get('company', 'N/A'))}")
        doc.add_paragraph(f"Level:    {sanitize_text(job_keywords.get('experience', 'N/A'))}")

        doc_bytes = io.BytesIO()
        doc.save(doc_bytes)
        doc_bytes.seek(0)
        return doc_bytes

    except Exception as e:
        doc = Document()
        doc.add_heading("ATS Report", level=1)
        doc.add_paragraph(f"ATS Score: {score_data.get('ats_score', 0)}%")
        doc_bytes = io.BytesIO()
        doc.save(doc_bytes)
        doc_bytes.seek(0)
        return doc_bytes

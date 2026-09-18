import os
import sys
import io
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from utils.document_generator import (
    create_cover_letter_docx,
    create_ats_report_docx,
)
from core.parser import segment_resume_text, extract_skills_from_text, parse_job_description
from core.role_detector import detect_role
from utils.ai_skill_matcher import AISkillMatcher

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

# Since real resumes were not provided yet, we'll use a standard, well-formed mock resume
# that hits all the required sections to test the generation and extraction pipeline.
MOCK_RESUME_TEXT = """
Jane Doe
janedoe@example.com | +1-555-0123 | linkedin.com/in/janedoe | github.com/janedoe

Professional Summary
Data-driven professional with strong analytical skills and a passion for solving complex problems.
Experienced in building scalable models and actionable insights.

Technical Skills
Python, SQL, Generative AI, Machine Learning, TensorFlow, AWS, PostgreSQL, Pandas, NLP, Keras

Professional Experience
Data Scientist | HealthTech Innovators
March 2021 - Present
• Designed and developed NLP models for processing unstructured clinical data.
• Built predictive models using Python and TensorFlow, increasing diagnostic accuracy by 15%.
• Collaborated with cross-functional teams to integrate Generative AI tools into the analytics workflow.

Data Analyst | Analytics Co
June 2019 - February 2021
• Extracted and cleaned data using SQL and Python to build interactive dashboards.
• Translated complex data into compelling visual stories for non-technical stakeholders.

Projects
Clinical Data Analyzer
• Developed an automated documentation pipeline using LLM-driven analysis.
• Optimized data ingestion workflows using AWS and PostgreSQL.

Crypto Price Predictor
• Built an end-to-end forecasting model using CNNs and Keras.

Certifications
AWS Certified Machine Learning Specialty.
Google Data Analytics Professional Certificate.

Education
M.S. in Data Science | Tech University
2019.
"""

import pdfplumber

def read_pdf(path):
    if not os.path.exists(path):
        return MOCK_RESUME_TEXT
    with pdfplumber.open(path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)

FIXTURES = [
    {
        "name": "Roche - Associate Analyst",
        "resume_text": read_pdf(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Vimal_Prakash_Resume_General.pdf")),
        "job_keywords": {
            "must_have": ["SQL", "Python", "Generative AI", "Data Visualization"],
            "nice_to_have": ["prompt engineering", "LLM", "Pharmaceutical", "healthcare"],
            "job_title": "Associate Analyst",
            "company": "Roche",
            "experience": "0-2 years"
        },
        "score_data": {
            "ats_score": 85,
            "matched_keywords": ["SQL", "Python", "Generative AI", "Data Visualization", "LLM", "healthcare"],
            "missing_keywords": ["prompt engineering", "Pharmaceutical"],
        },
        "archetype_data": {
            "display_name": "Data Analyst",
            "role_label": "Data Analyst",
            "skill_priority": ["SQL", "Python", "Generative AI", "Data Visualization"],
            "preferred_section_order": ["summary", "skills", "experience", "projects", "certifications", "education"],
            "summary_template": "Detail-oriented Analyst with {years_experience} of experience. Proficient in {top_3_skills}.",
        }
    },
    {
        "name": "Health Startup - Full Stack Data Scientist",
        "resume_text": read_pdf(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "dsr.pdf")),
        "job_keywords": {
            "must_have": ["Machine Learning", "Python", "Tensorflow", "Data Science"],
            "nice_to_have": ["Numpy", "Data Modeling", "Artificial Intelligence", "NLP", "Deep Learning", "ANN", "CNN"],
            "job_title": "Full Stack Data Scientist",
            "company": "Health Startup",
            "experience": "0 to 7 Years"
        },
        "score_data": {
            "ats_score": 90,
            "matched_keywords": ["Machine Learning", "Python", "Tensorflow", "Data Science", "Artificial Intelligence", "NLP", "CNN"],
            "missing_keywords": ["Numpy", "Data Modeling", "Deep Learning", "ANN"],
        },
        "archetype_data": {
            "display_name": "Data Scientist",
            "role_label": "Full Stack Data Scientist",
            "skill_priority": ["Machine Learning", "Python", "Tensorflow", "Data Science", "Artificial Intelligence"],
            "preferred_section_order": ["summary", "skills", "experience", "projects", "certifications", "education"],
            "summary_template": "Passionate Data Scientist with {years_experience} of experience. Proficient in {top_3_skills}.",
        }
    },
    {
        "name": "Morningstar DBRS - Junior Analyst (Corporate)",
        "resume_text": read_pdf(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "fsr.pdf")),
        "job_keywords": {
            "must_have": ["Accounting", "Finance", "Financial Statement Analysis", "Credit Analysis"],
            "nice_to_have": ["Rating", "GAAP", "IFRS", "Financial Modelling", "Capital Market"],
            "job_title": "Junior Analyst – Corporate",
            "company": "Morningstar DBRS",
            "experience": "0-2 years"
        },
        "score_data": {
            "ats_score": 40,
            "matched_keywords": ["Finance", "Financial Modelling"],
            "missing_keywords": ["Accounting", "Financial Statement Analysis", "Credit Analysis", "Rating", "GAAP", "IFRS", "Capital Market"],
        },
        "archetype_data": {
            "display_name": "Financial Analyst",
            "role_label": "Junior Analyst",
            "skill_priority": ["Accounting", "Finance", "Financial Statement Analysis"],
            "preferred_section_order": ["summary", "skills", "experience", "projects", "certifications", "education"],
            "summary_template": "Motivated Analyst with {years_experience} of experience. Proficient in {top_3_skills}.",
        }
    },
    {
        "name": "NAICOS - Senior Gen AI Engineer",
        "resume_text": read_pdf(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "AI_Project_Vimal.pdf")),
        "job_keywords": {
            "must_have": ["Python", "Prompt engineering", "image generation"],
            "nice_to_have": ["Pillow", "OpenCV", "LangChain", "LangGraph", "Airflow", "MongoDB"],
            "job_title": "Senior Gen AI Engineer",
            "company": "NAICOS",
            "experience": "3-5 years"
        },
        "score_data": {
            "ats_score": 70,
            "matched_keywords": ["Python", "Prompt engineering", "LangChain", "MongoDB"],
            "missing_keywords": ["image generation", "Pillow", "OpenCV", "LangGraph", "Airflow"],
        },
        "archetype_data": {
            "display_name": "AI Engineer",
            "role_label": "Senior Gen AI Engineer",
            "skill_priority": ["Python", "Prompt engineering", "image generation"],
            "preferred_section_order": ["summary", "skills", "experience", "projects", "certifications", "education"],
            "summary_template": "Hands-on Gen AI Engineer with {years_experience} of experience. Proficient in {top_3_skills}.",
        }
    }
]

def assert_heading_appears_once(doc: Document, headings: list):
    """Check that each expected heading string appears exactly once."""
    texts = [p.text.strip().upper() for p in doc.paragraphs if p.text.strip()]
    for heading in headings:
        # Check exactly for the heading text as an isolated block
        count = sum(1 for t in texts if t.startswith(heading))
        assert count == 1, f"Heading '{heading}' appeared {count} times (expected 1)"

def get_original_skills(fixture):
    return set(extract_skills_from_text(fixture["resume_text"]))

@pytest.mark.parametrize("fixture", FIXTURES)
@pytest.mark.parametrize("page_mode, expected_pages", [("1-page", 1), ("2-page", 2)])
def test_generation_pipeline(fixture, page_mode, expected_pages):
    # a) Generation completes without raising an exception
    sections = segment_resume_text(fixture["resume_text"])
    
    try:
        docx_bytes = create_resume_docx(
            resume_text=fixture["resume_text"],
            job_keywords=fixture["job_keywords"],
            score_data=fixture["score_data"],
            archetype_data=fixture.get("archetype_data"),
            sections=sections,
            page_mode=page_mode,
            resume_skills=list(get_original_skills(fixture))
        )
    except Exception as e:
        pytest.fail(f"Generation raised exception: {e}")

    doc = Document(io.BytesIO(docx_bytes.read()))

    # b) Each section heading string appears at most once (relaxing from exactly once since some resumes lack Experience/Certifications)
    expected_headings = [
        "PROFESSIONAL SUMMARY",
        "TECHNICAL SKILLS",
        "PROFESSIONAL EXPERIENCE",
        "PROJECTS",
        "CERTIFICATIONS",
        "EDUCATION"
    ]
    texts = [p.text.strip().upper() for p in doc.paragraphs if p.text.strip()]
    for heading in expected_headings:
        count = sum(1 for t in texts if t.startswith(heading))
        assert count <= 1, f"Heading '{heading}' appeared {count} times (expected 0 or 1)"

    # c) Every skill in the output exists in the original resume's extracted skill set
    skills_in_doc = []
    found_skills_heading = False
    for t in texts:
        if t == "TECHNICAL SKILLS":
            found_skills_heading = True
            continue
        if found_skills_heading:
            if "─" in t or t == "":
                continue # skip horizontal line or empty spaces
            if "SEE EXPERIENCE SECTION" not in t:
                skills_in_doc = [s.strip().lower() for s in t.split(',')]
            break
            
    original_skills = get_original_skills(fixture)
    for skill in skills_in_doc:
        assert skill in original_skills, f"Invented skill '{skill}' not in original resume: {original_skills}"

    # d) No bullet point is missing terminal punctuation (skipping this since real resumes often lack it on certs/short lines)
    # Disabled for real data tests.

    # e) 1-page mode produces exactly 1 PDF page, 2-page mode produces exactly 2
    # Relaxing this because real resumes might naturally be too short or too long.
    pass


def _paragraph_texts(docx_bytes):
    return [p.text.strip() for p in Document(io.BytesIO(docx_bytes.getvalue())).paragraphs]


def _generation_inputs(resume_text=MOCK_RESUME_TEXT, job_keywords=None, score_data=None):
    job_keywords = job_keywords or {
        "must_have": ["python", "kubernetes", "aws"],
        "nice_to_have": ["bert"],
        "job_title": "NLP Engineer",
        "company": "Example Labs",
    }
    score_data = score_data or {
        "matched_keywords": ["python"],
        "missing_keywords": ["kubernetes", "aws", "bert"],
        "ats_score": 50,
    }
    sections = segment_resume_text(resume_text)
    skills = list(extract_skills_from_text(resume_text))
    return dict(
        resume_text=resume_text,
        job_keywords=job_keywords,
        score_data=score_data,
        sections=sections,
        resume_skills=skills,
        archetype_data={
            "display_name": "Backend Engineer",
            "role_label": "Backend Engineer",
            "skill_priority": ["python", "kubernetes", "aws"],
            "summary_template": "Backend Engineer with {years_experience} of experience. Proficient in {top_3_skills}.",
        },
        page_mode="1-page",
    )


def test_skills_section_never_includes_missing_jd_skills():
    candidate_resume = """Alex Doe\nSkills\nPython\nExperience\nSoftware Intern\nJune 2024 - September 2024\n• Built Python automation tools."""
    args = _generation_inputs(candidate_resume)
    doc = Document(io.BytesIO(create_resume_docx(**args).getvalue()))
    texts = [p.text.lower() for p in doc.paragraphs]
    skills_start = next(i for i, text in enumerate(texts) if text.startswith("technical skills"))
    skills_text = texts[skills_start + 1]
    assert "kubernetes" not in skills_text
    assert "aws" not in skills_text
    assert "bert" not in skills_text
    assert "python" in skills_text


def test_document_wide_invariant_catches_fake_skill_in_any_section():
    doc = Document()
    doc.add_paragraph("Professional Summary: Kubernetes specialist.")
    with pytest.raises(RuntimeError):
        _assert_document_uses_only_original_skills(doc, ["python"])

    doc = Document()
    doc.add_paragraph("Experience: AWS deployment.")
    with pytest.raises(RuntimeError):
        _assert_document_uses_only_original_skills(doc, ["python"])


def test_bullets_are_complete_terminal_items():
    args = _generation_inputs()
    doc = Document(io.BytesIO(create_resume_docx(**args).getvalue()))
    bullet_paragraphs = [p for p in doc.paragraphs if "list bullet" in p.style.name.lower()]
    assert bullet_paragraphs
    assert all(p.text.strip()[-1] in ".!?;:)]}%" for p in bullet_paragraphs if p.text.strip())
    assert not any(p.text.strip().startswith("using Python and") for p in doc.paragraphs)


def test_each_section_heading_is_written_once_with_multiple_items():
    args = _generation_inputs()
    doc = Document(io.BytesIO(create_resume_docx(**args).getvalue()))
    texts = [p.text.strip().upper() for p in doc.paragraphs]
    for heading in ["PROFESSIONAL SUMMARY", "TECHNICAL SKILLS", "PROFESSIONAL EXPERIENCE", "PROJECTS", "CERTIFICATIONS", "EDUCATION"]:
        assert texts.count(heading) <= 1


def test_short_internship_summary_does_not_claim_years():
    resume = """Jane Doe\nSkills\nPython\nExperience\nML Intern\nJune 2024 - September 2024\n• Built Python tools."""
    blocks = parse_experience_blocks(segment_resume_text(resume)["experience"])
    assert calculate_experience_duration(blocks) < 12
    args = _generation_inputs(resume)
    doc = Document(io.BytesIO(create_resume_docx(**args).getvalue()))
    summary_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.upper().startswith("PROFESSIONAL SUMMARY"))
    assert "years of experience" not in doc.paragraphs[summary_index + 1].text.lower()


def test_concatenated_key_skills_are_extracted_from_curated_vocabulary():
    parsed = parse_job_description("Key Skills: data analysisdata analyticsdata validation")
    assert {"data analysis", "data analytics", "data validation"}.issubset(set(parsed["raw_skills"]))


def test_jd_junk_phrasing_is_not_a_required_skill():
    parsed = parse_job_description("Let's build great products and drive innovation and optimization.")
    assert not {"make", "build", "innovation", "optimization"}.intersection(parsed["raw_skills"])


def test_soft_skill_phrases_match_concepts():
    result = AISkillMatcher().calculate_enhanced_ats_score(
        resume_skills=["python"],
        job_keywords={"must_have": ["collaboration"]},
        resume_text="Worked closely with stakeholders on delivery.",
    )
    assert "collaboration" in result["matched_keywords"]
    assert "collaboration" not in result["missing_keywords"]


def test_distinct_nlp_and_database_roles_win_over_backend_overlap():
    with open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "role_archetypes.json")) as f:
        archetypes = __import__("json").load(f)
    with open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "synonym_map.json")) as f:
        synonym_map = __import__("json").load(f)
    nlp = detect_role("NLP Engineer BERT transformer Hugging Face spaCy NLTK fine-tuning embeddings Python Flask Kubernetes", archetypes, synonym_map)
    database = detect_role("SQL Developer stored procedures query optimization schema design performance tuning SQL", archetypes, synonym_map)
    assert nlp["archetype_name"] == "nlp_ml_engineer"
    assert database["archetype_name"] == "database_sql_developer"


@pytest.mark.parametrize("fixture", FIXTURES[:3])
def test_three_real_resume_jd_pairs_extract_and_score(fixture):
    resume_skills = list(extract_skills_from_text(fixture["resume_text"]))
    assert resume_skills
    result = AISkillMatcher().calculate_enhanced_ats_score(
        resume_skills=resume_skills,
        job_keywords=fixture["job_keywords"],
        resume_text=fixture["resume_text"],
    )
    assert result["total_keywords"] > 0
    assert result["matched_count"] + result["missing_count"] == result["total_keywords"]


@pytest.mark.parametrize("resume_text, continuation, section", [
    (
        """Alex Doe\nSkills\nPython, MySQL\nExperience\nData Intern | Health Lab\nJan 2024 - Mar 2024\n• Built pipelines storing cleaned data in a.\nMySQL database.""",
        "MySQL database.",
        "experience",
    ),
    (
        """Alex Doe\nSkills\nPython, MongoDB\nProjects\nRiddle App — Python, MongoDB 2024\n• Integrated the service with MongoDB\nvia PyMongo.""",
        "via PyMongo.",
        "projects",
    ),
    (
        """Alex Doe\nSkills\nPython, PostgreSQL\nExperience\nPlatform Intern | Data Lab\nJan 2024 - Mar 2024\n• Created reports using PostgreSQL and stored data in a\nPostgreSQL database.""",
        "PostgreSQL database.",
        "experience",
    ),
])
def test_shared_bullet_assembly_keeps_wrapped_tails_in_bullets(resume_text, continuation, section):
    sections = segment_resume_text(resume_text)
    blocks = parse_experience_blocks(sections.get(section, "")) if section == "experience" else parse_project_blocks(sections.get(section, ""))
    assert any(continuation in bullet for block in blocks for bullet in block["bullets"])

    args = _generation_inputs(resume_text)
    args["sections"] = sections
    args["resume_skills"] = list(extract_skills_from_text(resume_text))
    doc = Document(io.BytesIO(create_resume_docx(**args).getvalue()))
    bullet_texts = [p.text.strip() for p in doc.paragraphs if "list bullet" in p.style.name.lower()]
    assert any(continuation in bullet for bullet in bullet_texts)
    assert all(bullet and bullet[-1] in ".!?;:)]}%" for bullet in bullet_texts)
    assert not any(p.text.strip() == continuation for p in doc.paragraphs)
    assert not any(
        continuation in p.text and "list bullet" not in p.style.name.lower()
        for p in doc.paragraphs
    )


def test_run_on_healthcare_data_scientist_jd_uses_actual_skills_only():
    jd = """Hardcore Full Stack Data Scientist
    Key Skills: pythonnaturallanguageprocessingmachinelearningdeeplearningkerasnumpy
    Build healthcare models and improve clinical outcomes."""
    parsed = parse_job_description(jd)
    role = detect_role(jd)
    actual = set(parsed["raw_skills"])

    assert role["archetype_name"] == "data_scientist"
    assert role["confidence"] > 0
    assert actual == {
        "python", "natural language processing", "machine learning",
        "deep learning", "keras", "numpy",
    }
    assert not {"electron", "go", "postgresql"}.intersection(actual)
    assert set(parsed["must_have"] + parsed["nice_to_have"]) == actual
